from flask import Blueprint, render_template, redirect, url_for, flash, request, jsonify
from flask_login import login_required, current_user
from app.models import (
    db, Tenant, User, Restaurant, Table, Reservation, AIConversation,
    ROLE_ADMIN, ROLE_TENANT_SUPERUSER, ROLE_TENANT_USER, TENANT_ROLES
)
from functools import wraps
from werkzeug.security import generate_password_hash
import secrets

admin_bp = Blueprint('admin', __name__)


def generate_webhook_token():
    """Generate a unique webhook token for a restaurant"""
    while True:
        token = secrets.token_urlsafe(32)
        # Ensure token is unique
        existing = Restaurant.query.filter_by(webhook_token=token).first()
        if not existing:
            return token


# =============================================================================
# ACCESS CONTROL DECORATORS
# =============================================================================

def admin_only(f):
    """Decorator: Only system admin can access"""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not current_user.is_authenticated:
            return redirect(url_for('auth.login'))
        if not current_user.is_admin():
            flash('Access denied. System administrator privileges required.', 'error')
            return redirect(url_for('admin.dashboard'))
        return f(*args, **kwargs)
    return decorated_function


def tenant_access_required(f):
    """Decorator: User must have access to a tenant (admin or tenant user)"""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not current_user.is_authenticated:
            return redirect(url_for('auth.login'))
        # Admin has access to everything
        if current_user.is_admin():
            return f(*args, **kwargs)
        # Tenant users must have a tenant_id
        if not current_user.tenant_id:
            flash('Access denied. No tenant associated with your account.', 'error')
            return redirect(url_for('main.index'))
        return f(*args, **kwargs)
    return decorated_function


def tenant_superuser_required(f):
    """Decorator: User must be admin or tenant superuser"""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not current_user.is_authenticated:
            return redirect(url_for('auth.login'))
        if current_user.is_admin():
            return f(*args, **kwargs)
        if not current_user.is_tenant_superuser():
            flash('Access denied. Tenant administrator privileges required.', 'error')
            return redirect(url_for('admin.dashboard'))
        return f(*args, **kwargs)
    return decorated_function


# =============================================================================
# HELPER FUNCTIONS
# =============================================================================

def get_user_tenant():
    """Get the tenant for the current user (None for admin)"""
    if current_user.is_admin():
        return None
    return Tenant.query.get(current_user.tenant_id)


def get_accessible_restaurants():
    """Get restaurants accessible to current user"""
    if current_user.is_admin():
        return Restaurant.query.all()
    return Restaurant.query.filter_by(tenant_id=current_user.tenant_id).all()


def can_access_restaurant(restaurant):
    """Check if current user can access a restaurant"""
    if current_user.is_admin():
        return True
    return restaurant.tenant_id == current_user.tenant_id


def can_manage_restaurant(restaurant):
    """Check if current user can manage (edit) a restaurant"""
    if current_user.is_admin():
        return True
    return restaurant.tenant_id == current_user.tenant_id and current_user.is_tenant_superuser()


# =============================================================================
# DASHBOARD
# =============================================================================

@admin_bp.route('/')
@admin_bp.route('/dashboard')
@login_required
@tenant_access_required
def dashboard():
    """Admin dashboard - different views based on role"""
    if current_user.is_admin():
        # System admin sees all tenants overview
        tenants = Tenant.query.all()
        total_restaurants = Restaurant.query.count()
        total_reservations = Reservation.query.count()
        total_users = User.query.filter(User.role != ROLE_ADMIN).count()
        return render_template('admin/dashboard.html', 
                             tenants=tenants,
                             total_restaurants=total_restaurants,
                             total_reservations=total_reservations,
                             total_users=total_users,
                             is_admin=True)
    else:
        # Tenant users see their tenant's data
        tenant = Tenant.query.get(current_user.tenant_id)
        restaurants = Restaurant.query.filter_by(tenant_id=current_user.tenant_id).all()
        reservations = Reservation.query.join(Restaurant).filter(
            Restaurant.tenant_id == current_user.tenant_id
        ).order_by(Reservation.reservation_date.desc()).limit(10).all()
        tenant_users = User.query.filter_by(tenant_id=current_user.tenant_id).count()
        return render_template('admin/dashboard.html',
                             tenant=tenant,
                             restaurants=restaurants,
                             recent_reservations=reservations,
                             tenant_users=tenant_users,
                             is_admin=False)


# =============================================================================
# TENANT MANAGEMENT (System Admin Only)
# =============================================================================

@admin_bp.route('/tenants')
@login_required
@admin_only
def tenants():
    """List all tenants (system admin only)"""
    tenants = Tenant.query.all()
    return render_template('admin/tenants.html', tenants=tenants)


@admin_bp.route('/tenants/add', methods=['GET', 'POST'])
@login_required
@admin_only
def add_tenant():
    """Add new tenant (system admin only)"""
    if request.method == 'POST':
        tenant = Tenant(
            name=request.form.get('name'),
            email=request.form.get('email'),
            phone=request.form.get('phone'),
            subscription_plan=request.form.get('subscription_plan', 'starter'),
            subscription_status='active'
        )
        db.session.add(tenant)
        db.session.commit()
        flash('Client added successfully!', 'success')
        return redirect(url_for('admin.tenants'))
    
    return render_template('admin/tenant_form.html')


@admin_bp.route('/tenants/<int:id>')
@login_required
@admin_only
def tenant_detail(id):
    """View tenant details (system admin only)"""
    tenant = Tenant.query.get_or_404(id)
    return render_template('admin/tenant_detail.html', tenant=tenant)


@admin_bp.route('/tenants/<int:id>/edit', methods=['GET', 'POST'])
@login_required
@admin_only
def edit_tenant(id):
    """Edit tenant (system admin only)"""
    tenant = Tenant.query.get_or_404(id)
    
    if request.method == 'POST':
        tenant.name = request.form.get('name')
        tenant.email = request.form.get('email')
        tenant.phone = request.form.get('phone')
        tenant.subscription_plan = request.form.get('subscription_plan', 'starter')
        tenant.subscription_status = request.form.get('subscription_status', 'active')
        tenant.is_active = request.form.get('is_active') == 'on'
        db.session.commit()
        flash('Client updated successfully!', 'success')
        return redirect(url_for('admin.tenant_detail', id=id))
    
    return render_template('admin/tenant_form.html', tenant=tenant)


# =============================================================================
# USER MANAGEMENT
# =============================================================================

@admin_bp.route('/users')
@login_required
@tenant_access_required
def users():
    """List users - admin sees all, tenant superuser sees their tenant's users"""
    if current_user.is_admin():
        users = User.query.all()
        tenants = Tenant.query.all()
        return render_template('admin/users.html', users=users, tenants=tenants, is_admin=True)
    elif current_user.is_tenant_superuser():
        users = User.query.filter_by(tenant_id=current_user.tenant_id).all()
        return render_template('admin/users.html', users=users, is_admin=False)
    else:
        flash('Access denied. You do not have permission to manage users.', 'error')
        return redirect(url_for('admin.dashboard'))


@admin_bp.route('/users/add', methods=['GET', 'POST'])
@login_required
@tenant_superuser_required
def add_user():
    """Add new user"""
    # Get tenants for admin
    tenants = None
    if current_user.is_admin():
        tenants = Tenant.query.all()
    
    if request.method == 'POST':
        email = request.form.get('email')
        
        # Check if email already exists
        if User.query.filter_by(email=email).first():
            flash('A user with this email already exists.', 'error')
            return render_template('admin/user_form.html', tenants=tenants)
        
        # Determine tenant_id and role
        if current_user.is_admin():
            tenant_id = request.form.get('tenant_id')
            if tenant_id:
                tenant_id = int(tenant_id)
            else:
                tenant_id = None
            role = request.form.get('role', ROLE_TENANT_USER)
        else:
            tenant_id = current_user.tenant_id
            # Tenant superuser can only create tenant_user or tenant_superuser
            role = request.form.get('role', ROLE_TENANT_USER)
            if role not in TENANT_ROLES:
                role = ROLE_TENANT_USER
        
        user = User(
            email=email,
            first_name=request.form.get('first_name'),
            last_name=request.form.get('last_name'),
            role=role,
            tenant_id=tenant_id,
            is_active=True
        )
        user.set_password(request.form.get('password'))
        
        db.session.add(user)
        db.session.commit()
        flash('User added successfully!', 'success')
        return redirect(url_for('admin.users'))
    
    return render_template('admin/user_form.html', tenants=tenants)


@admin_bp.route('/users/<int:id>/edit', methods=['GET', 'POST'])
@login_required
@tenant_superuser_required
def edit_user(id):
    """Edit user"""
    user = User.query.get_or_404(id)
    
    # Check access
    if not current_user.is_admin():
        if user.tenant_id != current_user.tenant_id:
            flash('Access denied.', 'error')
            return redirect(url_for('admin.users'))
        # Tenant superuser cannot edit admin users
        if user.is_admin():
            flash('Access denied. Cannot edit system administrator.', 'error')
            return redirect(url_for('admin.users'))
    
    # Get tenants for admin
    tenants = None
    if current_user.is_admin():
        tenants = Tenant.query.all()
    
    if request.method == 'POST':
        user.first_name = request.form.get('first_name')
        user.last_name = request.form.get('last_name')
        user.is_active = request.form.get('is_active') == 'on'
        
        # Update password if provided
        new_password = request.form.get('password')
        if new_password:
            user.set_password(new_password)
        
        # Admin can change tenant and role
        if current_user.is_admin():
            tenant_id = request.form.get('tenant_id')
            user.tenant_id = int(tenant_id) if tenant_id else None
            user.role = request.form.get('role', user.role)
        else:
            # Tenant superuser can only change role within tenant roles
            role = request.form.get('role', user.role)
            if role in TENANT_ROLES:
                user.role = role
        
        db.session.commit()
        flash('User updated successfully!', 'success')
        return redirect(url_for('admin.users'))
    
    return render_template('admin/user_form.html', user=user, tenants=tenants)


@admin_bp.route('/users/<int:id>/delete', methods=['POST'])
@login_required
@tenant_superuser_required
def delete_user(id):
    """Delete user"""
    user = User.query.get_or_404(id)
    
    # Check access
    if not current_user.is_admin():
        if user.tenant_id != current_user.tenant_id:
            flash('Access denied.', 'error')
            return redirect(url_for('admin.users'))
        # Tenant superuser cannot delete admin users
        if user.is_admin():
            flash('Access denied. Cannot delete system administrator.', 'error')
            return redirect(url_for('admin.users'))
    
    # Cannot delete yourself
    if user.id == current_user.id:
        flash('You cannot delete your own account.', 'error')
        return redirect(url_for('admin.users'))
    
    db.session.delete(user)
    db.session.commit()
    flash(f'User {user.full_name} has been deleted.', 'success')
    return redirect(url_for('admin.users'))


# =============================================================================
# PASSWORD MANAGEMENT
# =============================================================================

@admin_bp.route('/users/<int:id>/change-password', methods=['GET', 'POST'])
@login_required
@tenant_superuser_required
def change_user_password(id):
    """Change user password - dedicated route for password changes"""
    user = User.query.get_or_404(id)
    
    # Check access
    if not current_user.is_admin():
        if user.tenant_id != current_user.tenant_id:
            flash('Access denied.', 'error')
            return redirect(url_for('admin.users'))
        # Tenant superuser cannot change admin passwords
        if user.is_admin():
            flash('Access denied. Cannot change system administrator password.', 'error')
            return redirect(url_for('admin.users'))
    
    if request.method == 'POST':
        new_password = request.form.get('new_password')
        confirm_password = request.form.get('confirm_password')
        
        # Validate passwords
        if not new_password:
            flash('Please enter a new password.', 'error')
            return render_template('admin/change_password.html', user=user)
        
        if len(new_password) < 8:
            flash('Password must be at least 8 characters long.', 'error')
            return render_template('admin/change_password.html', user=user)
        
        if new_password != confirm_password:
            flash('Passwords do not match.', 'error')
            return render_template('admin/change_password.html', user=user)
        
        # Update password
        user.set_password(new_password)
        db.session.commit()
        
        flash(f'Password for {user.full_name} ({user.email}) has been changed successfully!', 'success')
        return redirect(url_for('admin.users'))
    
    return render_template('admin/change_password.html', user=user)


# =============================================================================
# RESTAURANT MANAGEMENT
# =============================================================================

@admin_bp.route('/restaurants')
@login_required
@tenant_access_required
def restaurants():
    """List restaurants"""
    if current_user.is_admin():
        restaurants = Restaurant.query.all()
        tenants = Tenant.query.all()
        return render_template('admin/restaurants.html', restaurants=restaurants, tenants=tenants, is_admin=True)
    else:
        restaurants = Restaurant.query.filter_by(tenant_id=current_user.tenant_id).all()
        return render_template('admin/restaurants.html', restaurants=restaurants, is_admin=False)


@admin_bp.route('/restaurants/add', methods=['GET', 'POST'])
@login_required
@tenant_superuser_required
def add_restaurant():
    """Add new restaurant"""
    # Get tenants for admin to select from
    tenants = None
    if current_user.is_admin():
        tenants = Tenant.query.all()
    
    if request.method == 'POST':
        # Determine tenant_id
        if current_user.is_admin():
            tenant_id = request.form.get('tenant_id')
            if not tenant_id:
                flash('Please select a client for this restaurant.', 'error')
                return render_template('admin/restaurant_form.html', tenants=tenants)
            tenant_id = int(tenant_id)
        else:
            tenant_id = current_user.tenant_id
        
        # Generate unique webhook token for Chatwoot integration
        webhook_token = generate_webhook_token()
        
        restaurant = Restaurant(
            tenant_id=tenant_id,
            name=request.form.get('name'),
            address=request.form.get('address'),
            city=request.form.get('city'),
            state=request.form.get('state'),
            zip_code=request.form.get('zip_code'),
            phone=request.form.get('phone'),
            email=request.form.get('email'),
            cuisine_type=request.form.get('cuisine_type'),
            description=request.form.get('description'),
            timezone=request.form.get('timezone', 'UTC'),
            webhook_token=webhook_token
        )
        db.session.add(restaurant)
        db.session.commit()
        flash('Restaurant added successfully!', 'success')
        return redirect(url_for('admin.restaurants'))
    
    return render_template('admin/restaurant_form.html', tenants=tenants)


@admin_bp.route('/restaurants/<int:id>/edit', methods=['GET', 'POST'])
@login_required
@tenant_superuser_required
def edit_restaurant(id):
    """Edit restaurant"""
    restaurant = Restaurant.query.get_or_404(id)
    
    # Check access
    if not current_user.is_admin() and restaurant.tenant_id != current_user.tenant_id:
        flash('Access denied.', 'error')
        return redirect(url_for('admin.restaurants'))
    
    # Get tenants for admin
    tenants = None
    if current_user.is_admin():
        tenants = Tenant.query.all()
    
    if request.method == 'POST':
        # Update tenant_id if admin
        if current_user.is_admin():
            tenant_id = request.form.get('tenant_id')
            if tenant_id:
                restaurant.tenant_id = int(tenant_id)
        
        restaurant.name = request.form.get('name')
        restaurant.address = request.form.get('address')
        restaurant.city = request.form.get('city')
        restaurant.state = request.form.get('state')
        restaurant.zip_code = request.form.get('zip_code')
        restaurant.phone = request.form.get('phone')
        restaurant.email = request.form.get('email')
        restaurant.cuisine_type = request.form.get('cuisine_type')
        restaurant.description = request.form.get('description')
        restaurant.timezone = request.form.get('timezone', 'UTC')
        restaurant.is_active = request.form.get('is_active') == 'on'
        db.session.commit()
        flash('Restaurant updated successfully!', 'success')
        return redirect(url_for('admin.restaurants'))
    
    return render_template('admin/restaurant_form.html', restaurant=restaurant, tenants=tenants)


# =============================================================================
# RESTAURANT DETAIL & WIDGET
# =============================================================================

@admin_bp.route('/restaurants/<int:id>')
@login_required
@tenant_access_required
def restaurant_detail(id):
    """View restaurant details with AI assistant test"""
    restaurant = Restaurant.query.get_or_404(id)
    
    if not can_access_restaurant(restaurant):
        flash('Access denied.', 'error')
        return redirect(url_for('admin.restaurants'))
    
    # Get tables
    tables = Table.query.filter_by(restaurant_id=id).all()
    
    # Get recent reservations
    from datetime import date
    recent_reservations = Reservation.query.filter_by(
        restaurant_id=id
    ).order_by(Reservation.reservation_date.desc()).limit(10).all()
    
    # Get stats
    today = date.today()
    stats = {
        'tables': len(tables),
        'today_reservations': Reservation.query.filter_by(
            restaurant_id=id,
            reservation_date=today
        ).filter(Reservation.status.in_(['pending', 'confirmed'])).count(),
        'total_reservations': Reservation.query.filter_by(restaurant_id=id).count(),
        'ai_conversations': AIConversation.query.filter_by(restaurant_id=id).count()
    }
    
    return render_template('admin/restaurant_detail.html',
                           restaurant=restaurant,
                           tables=tables,
                           recent_reservations=recent_reservations,
                           stats=stats)


@admin_bp.route('/restaurants/<int:id>/widget')
@login_required
@tenant_access_required
def widget_code(id):
    """Get widget embed code for a restaurant"""
    restaurant = Restaurant.query.get_or_404(id)
    
    if not can_access_restaurant(restaurant):
        flash('Access denied.', 'error')
        return redirect(url_for('admin.restaurants'))
    
    # Get base URL
    base_url = request.host_url.rstrip('/')
    tenant = restaurant.tenant
    
    return render_template('admin/widget_code.html',
                           restaurant=restaurant,
                           tenant=tenant,
                           base_url=base_url,
                           is_paid=tenant.is_paid if tenant else False,
                           is_live=restaurant.is_live if hasattr(restaurant, 'is_live') else True)


# =============================================================================
# TABLE MANAGEMENT
# =============================================================================

@admin_bp.route('/restaurants/<int:restaurant_id>/tables')
@login_required
@tenant_access_required
def tables(restaurant_id):
    """List tables for a restaurant"""
    restaurant = Restaurant.query.get_or_404(restaurant_id)
    
    if not can_access_restaurant(restaurant):
        flash('Access denied.', 'error')
        return redirect(url_for('admin.restaurants'))
    
    tables = Table.query.filter_by(restaurant_id=restaurant_id).all()
    return render_template('admin/tables.html', restaurant=restaurant, tables=tables)


@admin_bp.route('/restaurants/<int:restaurant_id>/tables/add', methods=['GET', 'POST'])
@login_required
@tenant_superuser_required
def add_table(restaurant_id):
    """Add table to restaurant"""
    restaurant = Restaurant.query.get_or_404(restaurant_id)
    
    if not can_manage_restaurant(restaurant):
        flash('Access denied.', 'error')
        return redirect(url_for('admin.restaurants'))
    
    if request.method == 'POST':
        table = Table(
            restaurant_id=restaurant_id,
            table_number=request.form.get('table_number'),
            capacity=int(request.form.get('capacity')),
            location=request.form.get('location'),
            notes=request.form.get('notes')
        )
        db.session.add(table)
        db.session.commit()
        flash('Table added successfully!', 'success')
        return redirect(url_for('admin.tables', restaurant_id=restaurant_id))
    
    return render_template('admin/table_form.html', restaurant=restaurant)


@admin_bp.route('/restaurants/<int:restaurant_id>/tables/<int:id>/edit', methods=['GET', 'POST'])
@login_required
@tenant_superuser_required
def edit_table(restaurant_id, id):
    """Edit table"""
    restaurant = Restaurant.query.get_or_404(restaurant_id)
    table = Table.query.get_or_404(id)
    
    if not can_manage_restaurant(restaurant):
        flash('Access denied.', 'error')
        return redirect(url_for('admin.restaurants'))
    
    if request.method == 'POST':
        table.table_number = request.form.get('table_number')
        table.capacity = int(request.form.get('capacity'))
        table.location = request.form.get('location')
        table.notes = request.form.get('notes')
        table.is_available = request.form.get('is_available') == 'on'
        db.session.commit()
        flash('Table updated successfully!', 'success')
        return redirect(url_for('admin.tables', restaurant_id=restaurant_id))
    
    return render_template('admin/table_form.html', restaurant=restaurant, table=table)


# =============================================================================
# RESERVATION MANAGEMENT
# =============================================================================

@admin_bp.route('/reservations')
@login_required
@tenant_access_required
def reservations():
    """List reservations"""
    if current_user.is_admin():
        reservations = Reservation.query.order_by(Reservation.reservation_date.desc()).all()
    else:
        reservations = Reservation.query.join(Restaurant).filter(
            Restaurant.tenant_id == current_user.tenant_id
        ).order_by(Reservation.reservation_date.desc()).all()
    return render_template('admin/reservations.html', reservations=reservations)


@admin_bp.route('/reservations/<int:id>')
@login_required
@tenant_access_required
def reservation_detail(id):
    """View reservation details"""
    reservation = Reservation.query.get_or_404(id)
    
    if not current_user.is_admin() and reservation.restaurant.tenant_id != current_user.tenant_id:
        flash('Access denied.', 'error')
        return redirect(url_for('admin.reservations'))
    
    return render_template('admin/reservation_detail.html', reservation=reservation)


@admin_bp.route('/reservations/<int:id>/status', methods=['POST'])
@login_required
@tenant_access_required
def update_reservation_status(id):
    """Update reservation status"""
    reservation = Reservation.query.get_or_404(id)
    
    if not current_user.is_admin() and reservation.restaurant.tenant_id != current_user.tenant_id:
        flash('Access denied.', 'error')
        return redirect(url_for('admin.reservations'))
    
    new_status = request.form.get('status')
    if new_status in ['pending', 'confirmed', 'cancelled', 'completed', 'no_show']:
        reservation.status = new_status
        db.session.commit()
        flash(f'Reservation status updated to {new_status}.', 'success')
    
    return redirect(url_for('admin.reservation_detail', id=id))


# =============================================================================
# AI CONVERSATIONS LOG
# =============================================================================

@admin_bp.route('/conversations')
@login_required
@tenant_access_required
def conversations():
    """View AI conversation logs"""
    if current_user.is_admin():
        conversations = AIConversation.query.order_by(AIConversation.started_at.desc()).all()
    else:
        conversations = AIConversation.query.join(Restaurant).filter(
            Restaurant.tenant_id == current_user.tenant_id
        ).order_by(AIConversation.started_at.desc()).all()
    return render_template('admin/conversations.html', conversations=conversations)


# =============================================================================
# SETTINGS
# =============================================================================

@admin_bp.route('/settings', methods=['GET', 'POST'])
@login_required
@tenant_access_required
def settings():
    """Account and tenant settings"""
    if request.method == 'POST':
        # Update current user's profile
        current_user.first_name = request.form.get('first_name', current_user.first_name)
        current_user.last_name = request.form.get('last_name', current_user.last_name)
        
        # Update password if provided
        new_password = request.form.get('new_password')
        if new_password:
            current_password = request.form.get('current_password')
            if current_user.check_password(current_password):
                current_user.set_password(new_password)
                flash('Password updated successfully!', 'success')
            else:
                flash('Current password is incorrect.', 'error')
        
        db.session.commit()
        flash('Settings updated successfully!', 'success')
        return redirect(url_for('admin.settings'))
    
    tenant = None
    if not current_user.is_admin():
        tenant = Tenant.query.get(current_user.tenant_id)
    
    return render_template('admin/settings.html', tenant=tenant)


# =============================================================================
# TENANT INFO (For Tenant Users to View Their Tenant)
# =============================================================================

@admin_bp.route('/my-organization')
@login_required
@tenant_access_required
def my_organization():
    """View current user's organization/tenant info"""
    if current_user.is_admin():
        flash('System administrators do not belong to a specific organization.', 'info')
        return redirect(url_for('admin.tenants'))
    
    tenant = Tenant.query.get_or_404(current_user.tenant_id)
    return render_template('admin/my_organization.html', tenant=tenant)


@admin_bp.route('/my-organization/edit', methods=['GET', 'POST'])
@login_required
@tenant_superuser_required
def edit_my_organization():
    """Edit current user's organization/tenant info (tenant superuser only)"""
    if current_user.is_admin():
        flash('System administrators do not belong to a specific organization.', 'info')
        return redirect(url_for('admin.tenants'))
    
    tenant = Tenant.query.get_or_404(current_user.tenant_id)
    
    if request.method == 'POST':
        tenant.name = request.form.get('name', tenant.name)
        tenant.email = request.form.get('email', tenant.email)
        tenant.phone = request.form.get('phone', tenant.phone)
        db.session.commit()
        flash('Organization updated successfully!', 'success')
        return redirect(url_for('admin.my_organization'))
    
    return render_template('admin/edit_organization.html', tenant=tenant)


# =============================================================================
# STAFF ASSISTANT
# =============================================================================

@admin_bp.route('/staff-assistant')
@login_required
@tenant_access_required
def staff_assistant():
    """Staff Assistant page for managing reservations via AI chat"""
    # Get the restaurant for the current user
    if current_user.is_admin():
        # Admin can select any restaurant
        restaurants = Restaurant.query.all()
        restaurant = restaurants[0] if restaurants else None
    else:
        # Tenant user sees their tenant's restaurants
        restaurants = Restaurant.query.filter_by(tenant_id=current_user.tenant_id).all()
        restaurant = restaurants[0] if restaurants else None
    
    if not restaurant:
        flash('No restaurant found. Please create a restaurant first.', 'warning')
        return redirect(url_for('admin.restaurants'))
    
    return render_template('admin/staff_assistant.html', 
                         restaurant=restaurant, 
                         restaurants=restaurants)


# =============================================================================
# KNOWLEDGE BASE MANAGEMENT
# =============================================================================

@admin_bp.route('/restaurants/<int:id>/knowledge-base', methods=['GET', 'POST'])
@login_required
@tenant_superuser_required
def knowledge_base(id):
    """View and edit restaurant knowledge base"""
    restaurant = Restaurant.query.get_or_404(id)
    
    if not can_access_restaurant(restaurant):
        flash('Access denied.', 'error')
        return redirect(url_for('admin.restaurants'))
    
    if request.method == 'POST':
        from datetime import datetime
        
        # Check if file upload or text edit
        if 'knowledge_file' in request.files:
            file = request.files['knowledge_file']
            if file and file.filename:
                # Read the markdown file content
                content = file.read().decode('utf-8')
                restaurant.knowledge_base = content
                restaurant.knowledge_base_updated_at = datetime.utcnow()
                db.session.commit()
                flash('Knowledge base uploaded successfully!', 'success')
        else:
            # Text edit
            content = request.form.get('knowledge_base', '')
            restaurant.knowledge_base = content
            restaurant.knowledge_base_updated_at = datetime.utcnow()
            db.session.commit()
            flash('Knowledge base updated successfully!', 'success')
        
        return redirect(url_for('admin.knowledge_base', id=id))
    
    # Generate sample template if no knowledge base exists
    sample_template = """# Restaurant Knowledge Base

## About Us
[Describe your restaurant, its history, and what makes it special]

## Location & Contact
- **Address:** [Your full address]
- **Phone:** [Your phone number]
- **Email:** [Your email]
- **Parking:** [Parking information]

## Hours of Operation
| Day | Hours |
|-----|-------|
| Monday | 11:00 AM - 10:00 PM |
| Tuesday | 11:00 AM - 10:00 PM |
| Wednesday | 11:00 AM - 10:00 PM |
| Thursday | 11:00 AM - 10:00 PM |
| Friday | 11:00 AM - 11:00 PM |
| Saturday | 10:00 AM - 11:00 PM |
| Sunday | 10:00 AM - 9:00 PM |

## Menu Highlights

### Appetizers
- [Dish name] - $XX - [Brief description]
- [Dish name] - $XX - [Brief description]

### Main Courses
- [Dish name] - $XX - [Brief description]
- [Dish name] - $XX - [Brief description]

### Desserts
- [Dish name] - $XX - [Brief description]

### Beverages
- [List your signature drinks, wine selection, etc.]

## Dietary Options
- **Vegetarian:** [List vegetarian options]
- **Vegan:** [List vegan options]
- **Gluten-Free:** [List gluten-free options]
- **Allergen Info:** [Note about allergen accommodations]

## Reservations Policy
- Reservations recommended for parties of [X] or more
- Cancellation policy: [Your policy]
- Large party accommodations: [Details]

## Private Events
[Information about private dining, event spaces, catering]

## Frequently Asked Questions

### Do you have outdoor seating?
[Your answer]

### Is there a dress code?
[Your answer]

### Do you accommodate dietary restrictions?
[Your answer]

### Is the restaurant wheelchair accessible?
[Your answer]

### Do you offer takeout or delivery?
[Your answer]

## Special Features
- [Live music on weekends]
- [Happy hour specials]
- [Chef's tasting menu]
- [Wine pairing dinners]
"""
    
    return render_template('admin/knowledge_base.html', 
                         restaurant=restaurant,
                         sample_template=sample_template)


@admin_bp.route('/restaurants/<int:id>/knowledge-base/download')
@login_required
@tenant_access_required
def download_knowledge_base(id):
    """Download the knowledge base as a markdown file"""
    restaurant = Restaurant.query.get_or_404(id)
    
    if not can_access_restaurant(restaurant):
        flash('Access denied.', 'error')
        return redirect(url_for('admin.restaurants'))
    
    if not restaurant.knowledge_base:
        flash('No knowledge base to download.', 'warning')
        return redirect(url_for('admin.knowledge_base', id=id))
    
    from flask import Response
    
    # Create filename
    filename = f"{restaurant.name.lower().replace(' ', '_')}_knowledge_base.md"
    
    return Response(
        restaurant.knowledge_base,
        mimetype='text/markdown',
        headers={'Content-Disposition': f'attachment; filename={filename}'}
    )



# =============================================================================
# CHATWOOT INTEGRATION
# =============================================================================

@admin_bp.route('/restaurants/<int:id>/chatwoot', methods=['GET'])
@login_required
@tenant_superuser_required
def chatwoot_settings(id):
    """View Chatwoot integration settings for a restaurant"""
    restaurant = Restaurant.query.get_or_404(id)
    
    if not can_access_restaurant(restaurant):
        flash('Access denied.', 'error')
        return redirect(url_for('admin.restaurants'))
    
    # Generate webhook token if not exists (for existing restaurants)
    if not restaurant.webhook_token:
        restaurant.webhook_token = generate_webhook_token()
        db.session.commit()
    
    # Build the webhook URL
    webhook_url = f"{request.url_root.rstrip('/')}/api/webhook/chatwoot/{restaurant.webhook_token}"
    
    return render_template('admin/chatwoot_settings.html', 
                         restaurant=restaurant,
                         webhook_url=webhook_url)


@admin_bp.route('/restaurants/<int:id>/chatwoot/save', methods=['POST'])
@login_required
@tenant_superuser_required
def save_chatwoot_settings(id):
    """Save Chatwoot configuration for a restaurant"""
    restaurant = Restaurant.query.get_or_404(id)
    
    if not can_access_restaurant(restaurant):
        flash('Access denied.', 'error')
        return redirect(url_for('admin.restaurants'))
    
    # Update Chatwoot settings
    restaurant.chatwoot_base_url = request.form.get('chatwoot_base_url', '').strip()
    restaurant.chatwoot_account_id = request.form.get('chatwoot_account_id', '').strip()
    restaurant.chatwoot_inbox_id = request.form.get('chatwoot_inbox_id', '').strip()
    
    # Only update API key if provided (don't clear existing)
    api_key = request.form.get('chatwoot_api_key', '').strip()
    if api_key:
        restaurant.chatwoot_api_key = api_key
    
    db.session.commit()
    flash('Chatwoot settings saved successfully!', 'success')
    
    return redirect(url_for('admin.chatwoot_settings', id=id))


@admin_bp.route('/restaurants/<int:id>/chatwoot/regenerate-token', methods=['POST'])
@login_required
@tenant_superuser_required
def regenerate_webhook_token(id):
    """Regenerate the webhook token for a restaurant"""
    restaurant = Restaurant.query.get_or_404(id)
    
    if not can_access_restaurant(restaurant):
        flash('Access denied.', 'error')
        return redirect(url_for('admin.restaurants'))
    
    # Generate new token
    restaurant.webhook_token = generate_webhook_token()
    db.session.commit()
    
    flash('Webhook token regenerated. Please update the webhook URL in Chatwoot.', 'warning')
    
    return redirect(url_for('admin.chatwoot_settings', id=id))



# =============================================================================
# FLOOR PLAN MANAGEMENT
# =============================================================================

@admin_bp.route('/restaurants/<int:restaurant_id>/floor-plan')
@login_required
@tenant_access_required
def floor_plan_editor(restaurant_id):
    """Display the floor plan editor for a restaurant"""
    restaurant = Restaurant.query.get_or_404(restaurant_id)
    
    if not can_access_restaurant(restaurant):
        flash('Access denied.', 'error')
        return redirect(url_for('admin.restaurants'))
    
    # Try to get existing floor plan
    floor_plan = None
    try:
        from app.models import FloorPlan
        floor_plan = FloorPlan.query.filter_by(
            restaurant_id=restaurant_id,
            is_active=True
        ).first()
    except Exception:
        # FloorPlan model may not exist yet
        pass
    
    return render_template('admin/floor_plan_editor.html',
                         restaurant=restaurant,
                         floor_plan=floor_plan)


@admin_bp.route('/restaurants/<int:restaurant_id>/floor-plan/save', methods=['POST'])
@login_required
@tenant_superuser_required
def save_floor_plan(restaurant_id):
    """Save floor plan configuration"""
    from pydantic import ValidationError
    from datetime import datetime
    
    restaurant = Restaurant.query.get_or_404(restaurant_id)
    
    if not can_manage_restaurant(restaurant):
        return jsonify({'success': False, 'error': 'Access denied'}), 403
    
    try:
        data = request.get_json()
        
        # Validate with Pydantic
        from app.routes.floor_plan_routes import FloorPlanInput
        floor_plan_data = FloorPlanInput(**data)
        
        # Import models
        from app.models import FloorPlan, TableConfig, FloorCell
        
        # Get or create floor plan
        floor_plan = FloorPlan.query.filter_by(
            restaurant_id=restaurant_id,
            is_active=True
        ).first()
        
        if not floor_plan:
            floor_plan = FloorPlan(restaurant_id=restaurant_id)
            db.session.add(floor_plan)
        
        # Update floor plan properties
        floor_plan.name = floor_plan_data.name
        floor_plan.grid_rows = floor_plan_data.grid_rows
        floor_plan.grid_cols = floor_plan_data.grid_cols
        floor_plan.cell_size = floor_plan_data.cell_size
        floor_plan.floor_color = floor_plan_data.floor_color
        floor_plan.updated_at = datetime.utcnow()
        
        db.session.flush()  # Get floor_plan.id
        
        # Clear existing tables and floor cells
        TableConfig.query.filter_by(floor_plan_id=floor_plan.id).delete()
        FloorCell.query.filter_by(floor_plan_id=floor_plan.id).delete()
        
        # Add tables
        for table_data in floor_plan_data.tables:
            table = TableConfig(
                floor_plan_id=floor_plan.id,
                table_id=table_data.table_id,
                table_name=table_data.table_name,
                seats=table_data.seats,
                shape=table_data.shape,
                width=table_data.width,
                height=table_data.height,
                pos_x=table_data.pos_x,
                pos_y=table_data.pos_y,
                table_type=table_data.table_type,
                is_active=table_data.is_active,
                min_guests=table_data.min_guests,
                notes=table_data.notes
            )
            db.session.add(table)
        
        # Add floor cells
        for cell_data in floor_plan_data.floor_cells:
            cell = FloorCell(
                floor_plan_id=floor_plan.id,
                pos_x=cell_data.pos_x,
                pos_y=cell_data.pos_y,
                cell_type=cell_data.cell_type,
                color=cell_data.color
            )
            db.session.add(cell)
        
        # =============================================================
        # SYNC TO MAIN TABLES TABLE
        # This ensures the restaurant view shows correct table counts
        # =============================================================
        
        # Get existing tables for this restaurant (by table_number)
        existing_tables = {t.table_number: t for t in Table.query.filter_by(restaurant_id=restaurant_id).all()}
        synced_table_ids = set()
        
        for table_data in floor_plan_data.tables:
            table_number = table_data.table_id  # e.g., "T1", "T2"
            
            if table_number in existing_tables:
                # Update existing table
                table = existing_tables[table_number]
                table.capacity = table_data.seats
                table.is_active = table_data.is_active
                table.notes = table_data.notes or table.notes
                # Set location based on table type
                if table_data.table_type:
                    table.location = table_data.table_type
            else:
                # Create new table
                table = Table(
                    restaurant_id=restaurant_id,
                    table_number=table_number,
                    capacity=table_data.seats,
                    location=table_data.table_type or 'indoor',
                    is_available=True,
                    is_active=table_data.is_active,
                    notes=table_data.notes
                )
                db.session.add(table)
            
            synced_table_ids.add(table_number)
        
        # Optionally: Deactivate tables that are no longer in the floor plan
        # (but don't delete them to preserve reservation history)
        for table_number, table in existing_tables.items():
            if table_number not in synced_table_ids:
                # Table was removed from floor plan - mark as inactive
                table.is_active = False
        
        db.session.commit()
        
        # Count synced tables for response
        total_tables = len(floor_plan_data.tables)
        total_seats = sum(t.seats for t in floor_plan_data.tables)
        
        return jsonify({
            'success': True,
            'floor_plan_id': floor_plan.id,
            'message': f'Floor plan saved successfully. {total_tables} tables with {total_seats} total seats synced.',
            'tables_synced': total_tables,
            'total_seats': total_seats
        })
        
    except ValidationError as e:
        return jsonify({
            'success': False,
            'error': 'Validation error',
            'details': e.errors()
        }), 400
    except Exception as e:
        db.session.rollback()
        import traceback
        traceback.print_exc()
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500


@admin_bp.route('/restaurants/<int:restaurant_id>/floor-plan/data')
@login_required
@tenant_access_required
def get_floor_plan_data(restaurant_id):
    """Get floor plan data as JSON"""
    restaurant = Restaurant.query.get_or_404(restaurant_id)
    
    if not can_access_restaurant(restaurant):
        return jsonify({'success': False, 'error': 'Access denied'}), 403
    
    try:
        from app.models import FloorPlan
        floor_plan = FloorPlan.query.filter_by(
            restaurant_id=restaurant_id,
            is_active=True
        ).first()
        
        if not floor_plan:
            return jsonify({
                'success': True,
                'floor_plan': None
            })
        
        return jsonify({
            'success': True,
            'floor_plan': floor_plan.to_dict()
        })
    except Exception as e:
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500


@admin_bp.route('/restaurants/<int:restaurant_id>/floor-plan/tables')
@login_required
@tenant_access_required
def get_floor_plan_tables(restaurant_id):
    """Get all tables for a restaurant's floor plan"""
    restaurant = Restaurant.query.get_or_404(restaurant_id)
    
    if not can_access_restaurant(restaurant):
        return jsonify({'success': False, 'error': 'Access denied'}), 403
    
    try:
        from app.models import FloorPlan, TableConfig
        
        floor_plan = FloorPlan.query.filter_by(
            restaurant_id=restaurant_id,
            is_active=True
        ).first()
        
        if not floor_plan:
            return jsonify({'success': True, 'tables': []})
        
        tables = TableConfig.query.filter_by(
            floor_plan_id=floor_plan.id,
            is_active=True
        ).all()
        
        return jsonify({
            'success': True,
            'tables': [t.to_dict() for t in tables]
        })
    except Exception as e:
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500


@admin_bp.route('/restaurants/<int:restaurant_id>/floor-plan/upload-excel', methods=['POST'])
@login_required
@tenant_superuser_required
def upload_excel_layout(restaurant_id):
    """Upload Excel file to parse table layout"""
    import openpyxl
    from io import BytesIO
    
    restaurant = Restaurant.query.get_or_404(restaurant_id)
    
    if not can_manage_restaurant(restaurant):
        return jsonify({'success': False, 'error': 'Access denied'}), 403
    
    if 'file' not in request.files:
        return jsonify({'success': False, 'error': 'No file uploaded'}), 400
    
    file = request.files['file']
    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({'success': False, 'error': 'Invalid file type. Please upload an Excel file (.xlsx or .xls)'}), 400
    
    try:
        wb = openpyxl.load_workbook(BytesIO(file.read()))
        sheet = wb.active
        
        # Parse the Excel layout
        range_ref = sheet.dimensions
        if not range_ref or range_ref == 'A1:A1':
            return jsonify({'success': False, 'error': 'Empty spreadsheet'}), 400
        
        from openpyxl.utils import range_boundaries
        min_col, min_row, max_col, max_row = range_boundaries(range_ref)
        
        rows = max_row - min_row + 1
        cols = max_col - min_col + 1
        
        green_cells = []
        dark_cells = []
        
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                cell = sheet.cell(row=row, column=col)
                fill = cell.fill
                
                grid_row = row - min_row
                grid_col = col - min_col
                
                if fill and fill.fgColor:
                    color_type = fill.fgColor.type
                    if color_type == 'theme':
                        theme = fill.fgColor.theme
                        if theme == 9:  # Green = table
                            green_cells.append({'row': grid_row, 'col': grid_col})
                        elif theme == 1:  # Dark = floor
                            dark_cells.append({'row': grid_row, 'col': grid_col})
        
        # Group green cells into tables
        tables = _find_connected_cells(green_cells)
        
        parsed_tables = []
        for i, table_cells in enumerate(tables, 1):
            min_r = min(c['row'] for c in table_cells)
            max_r = max(c['row'] for c in table_cells)
            min_c = min(c['col'] for c in table_cells)
            max_c = max(c['col'] for c in table_cells)
            
            width = max_c - min_c + 1
            height = max_r - min_r + 1
            seats = len(table_cells)
            
            parsed_tables.append({
                'table_id': f'T{i}',
                'seats': seats,
                'width': width,
                'height': height,
                'pos_x': min_c,
                'pos_y': min_r,
                'shape': 'square' if width == height else 'rectangle'
            })
        
        return jsonify({
            'success': True,
            'grid_rows': rows,
            'grid_cols': cols,
            'tables': parsed_tables,
            'floor_cells': dark_cells
        })
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)}), 500


def _find_connected_cells(cells):
    """Find groups of connected cells (tables)"""
    if not cells:
        return []
    
    cell_set = set((c['row'], c['col']) for c in cells)
    groups = []
    
    while cell_set:
        start = cell_set.pop()
        group = [{'row': start[0], 'col': start[1]}]
        queue = [start]
        
        while queue:
            current = queue.pop(0)
            row, col = current
            
            neighbors = [
                (row-1, col), (row+1, col),
                (row, col-1), (row, col+1)
            ]
            
            for n in neighbors:
                if n in cell_set:
                    cell_set.remove(n)
                    group.append({'row': n[0], 'col': n[1]})
                    queue.append(n)
        
        groups.append(group)
    
    return groups
